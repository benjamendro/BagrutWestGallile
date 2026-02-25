#!/usr/bin/env python3
"""
Generate dynamic bagrut dashboards for East and West Galilee.
Each dashboard lets users select a subject and units level,
then dynamically updates all analytics and charts.
Also generates raw XLSX files for data validation.
"""
import openpyxl
import json
import os

BASE_DIR = '/home/user/BagrutWestGallile'

XLSX_FILES = [
    (os.path.join(BASE_DIR, 'BAGRUT.xlsx'), '2021-2023'),
    (os.path.join(BASE_DIR, 'BAGRUT תשפד.xlsx'), '2024'),
]

WEST_GALILEE_AUTHORITIES = {
    'עכו', 'נהריה', 'מעלות-תרשיחא', "בית ג'ן", 'חורפיש',
    'כפר ורדים', 'ירכא', 'אבו סנאן', 'כפר יאסיף', 'כסרא-סמיע',
    "ג'וליס", 'מעיליא', 'מטה אשר', 'מעלה יוסף', "ג'דיידה-מכר",
}

EAST_GALILEE_AUTHORITIES = {
    'חצור הגלילית', 'קרית שמונה', 'קצרין', 'גולן', 'צפת',
    'הגליל העליון', "מג'דל שמס", 'בוקעאתה', 'מסעדה', 'גוש חלב',
}

WEST_ALL_SCHOOLS = [
    'אבו-סלים סלמאן אלשיך', 'אולפנת צביה מעלות', 'אופק', 'אורט אולפנת הראל',
    'אורט מדעים ואומנויות', 'אורט מח"ט פסגות', 'אורט שחקים נהריה', 'אורט שלומי',
    'אמירים כפר ורדים', 'בי"ס מקיף השלום דנון', 'בית יעקב עכו', 'דרוזי מדעים ומנהיגות',
    "דרכא ג'וליס", "חט\"ע מכר-ג'דידה", 'טכנולוגי נעמ"ת', 'טרה סנטה',
    'ישיבה תיכונית', 'ישיבת אביר יעקב', 'כ"ג יורדי הסירה', 'כברי חט"ע',
    'מח"ט אורט מרום עכו', 'מנו"ף', 'מעלות רבקה', 'מקיף אורט מעלות',
    "מקיף אחווה ירכא", "מקיף בית ג'ן חט\"ב המ", 'מקיף גדידה', 'מקיף גליל מערבי',
    'מקיף דתי אמי"ת קנדי', 'מקיף סולם צור', 'מקיף ע"ש יני', 'מקיף ערבי',
    'מקיף שש - שנתי', 'מקיף תרשיחא', 'מרכז חינוך טכנולוגי', 'נוטרדאם',
    'ניסויי על איזורי תפן', 'נירים', 'עמל כסרא סמיע', 'שמריהו בירן',
]

EAST_ALL_SCHOOLS = [
    'שש שנתי בוקעאתא', 'אדם ואדמה גולן', 'תיכון פה הזהב', 'אברט חצור הגלילית',
    'קרית החינוך אמי"ת', 'מקיף טובא זנגריה', "עיוני מג'דל שמס", 'תיכון מסעדה',
    'ברנקו וייס מרום גליל', 'תיכון סלמאן חטיב', 'ישיבת תיכונית אמי"ת', 'עמל בגין צפת',
    'בית חנה', 'אולפנא אמי"ת', 'אולפנת קצרין', 'אורט דנציגר ק. שמונה',
    'המתמיד קרית שמונה', 'ברנקו וייס ק.שמונה', 'אולפנית קרית שמונה',
    'ישיבה תיכונית חיספין', 'ישיבת אלוני הבשן', 'רגבים בגולן',
    'ברנקו וייס גולן', 'כפר הנוער איילת השחר', 'אולפנת רגבים אופק', 'תיכון תמר',
    'מקיף עינות הירדן', 'מקיף עמק החולה', 'מקיף הר וגיא', 'מקיף אנה פרנק',
    'ישיבת בנע מירון', 'יב"ע בר יוחאי', 'דרכא נופי גולן',
]


def strip(val):
    if isinstance(val, str):
        return val.strip()
    return val


def read_school_data(authorities):
    """Read all school-level data from sheet 3, filtered by a set of authorities."""
    data = []
    for fpath, label in XLSX_FILES:
        if not os.path.exists(fpath):
            print(f"  File not found: {fpath}")
            continue
        print(f"  Reading {os.path.basename(fpath)} ({label})...")
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        ws = wb["סע' 3"]
        header_found = False
        for row in ws.iter_rows(values_only=True):
            rl = list(row)
            if not header_found:
                rt = ' '.join(str(c) for c in rl if c is not None)
                if 'ממוצע ציון סופי' in rt and 'מספר נבחנים' in rt:
                    header_found = True
                continue
            if rl[0] is None:
                continue
            auth = strip(rl[5])
            if auth not in authorities:
                continue
            data.append({
                'school': strip(rl[6]),
                'authority': auth,
                'supervision': strip(rl[8]),
                'sector': strip(rl[9]),
                'year': rl[3],
                'subject': strip(rl[4]),
                'units': rl[2],
                'examinees': rl[1],
                'score': rl[0],
            })
        wb.close()
    print(f"  Total rows for region: {len(data)}")
    return data


def read_national_data():
    """Read national-level data from sheet 1 (per subject+units+year)."""
    data = []
    for fpath, label in XLSX_FILES:
        if not os.path.exists(fpath):
            continue
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        ws = wb["סע' 1 "]
        header_found = False
        for row in ws.iter_rows(values_only=True):
            rl = list(row)
            if not header_found:
                rt = ' '.join(str(c) for c in rl if c is not None)
                if 'ממוצע ציון סופי' in rt and 'מספר נבחנים' in rt and 'יח"ל' in rt:
                    header_found = True
                continue
            if rl[0] is None:
                continue
            data.append({
                'subject': strip(rl[4]),
                'units': rl[2],
                'year': rl[3],
                'examinees': rl[1],
                'score': rl[0],
            })
        wb.close()
    return data


def generate_raw_xlsx(data, output_path):
    """Generate a raw data XLSX file for validation."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Raw Data'
    headers = ['שם מוסד', 'רשות מקומית', 'פיקוח', 'מגזר', 'מחזור סיום', 'מקצוע', 'יח"ל', 'מספר נבחנים', 'ממוצע ציון סופי']
    ws.append(headers)
    for d in sorted(data, key=lambda x: (x['subject'], x['units'], x['school'], x['year'])):
        ws.append([
            d['school'], d['authority'], d['supervision'], d['sector'],
            d['year'], d['subject'], d['units'], d['examinees'], d['score'],
        ])
    # Auto-filter
    ws.auto_filter.ref = f"A1:I{len(data)+1}"
    wb.save(output_path)
    print(f"  Raw XLSX saved: {output_path} ({len(data)} rows)")


def generate_html(region_heb, data, all_schools, national_data, output_path):
    """Generate an HTML dashboard with embedded data and dynamic JS."""
    json_data = json.dumps(data, ensure_ascii=False)
    json_schools = json.dumps(all_schools, ensure_ascii=False)
    json_national = json.dumps(national_data, ensure_ascii=False)

    html = HTML_TEMPLATE.replace('__RAW_DATA__', json_data)
    html = html.replace('__ALL_SCHOOLS__', json_schools)
    html = html.replace('__NATIONAL_DATA__', json_national)
    html = html.replace('__REGION_HEB__', region_heb)

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"  HTML saved: {output_path}")


HTML_TEMPLATE = r'''<!DOCTYPE html>
<html lang="he" dir="rtl">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>ניתוח מגמות בבגרות - __REGION_HEB__</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <script src="https://cdn.jsdelivr.net/npm/chartjs-plugin-annotation@3.0.1/dist/chartjs-plugin-annotation.min.js"></script>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Assistant:wght@400;600;700&display=swap" rel="stylesheet">
    <style>
        body { font-family: 'Assistant', sans-serif; }
        .card-container {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(320px, 1fr));
            gap: 1.5rem;
        }
        @media (min-width: 1024px) {
            .jewish-grid-container { grid-column: span 2 / span 2; }
            .jewish-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 1.5rem; }
        }
        .analysis-card {
            background-color: white;
            padding: 1.5rem;
            border-radius: 0.75rem;
            box-shadow: 0 4px 6px -1px rgb(0 0 0 / 0.1), 0 2px 4px -2px rgb(0 0 0 / 0.1);
        }
        select {
            padding: 0.5rem 1rem;
            border: 2px solid #cbd5e1;
            border-radius: 0.5rem;
            font-size: 1rem;
            font-family: 'Assistant', sans-serif;
            background: white;
            cursor: pointer;
            min-width: 200px;
        }
        select:focus { outline: none; border-color: #3b82f6; }
        optgroup { font-weight: 700; color: #1e293b; font-size: 1rem; }
        optgroup option { font-weight: 400; color: #334155; padding-right: 0.5rem; }
        .selector-group { display: flex; flex-wrap: wrap; justify-content: center; gap: 1rem; align-items: center; margin-top: 1.5rem; }
        .selector-label { font-weight: 600; color: #475569; }
        #no-data-msg { display: none; }
    </style>
</head>
<body class="bg-slate-50 text-slate-800">
    <div class="container mx-auto p-4 md:p-8">
        <header class="text-center mb-10">
            <img src="לוגו מרכז הידע.png" alt="לוגו מרכז הידע האזורי" class="mx-auto mb-6" style="max-height: 80px;">
            <h1 class="text-3xl md:text-4xl font-bold text-slate-900">ניתוח מגמות בבגרות - __REGION_HEB__</h1>
            <p class="text-lg text-slate-600 mt-2">סקירה כללית לפי מגזרים ומקצועות לשנים 2021-2024</p>

            <div class="selector-group">
                <div>
                    <span class="selector-label">מקצוע:</span>
                    <select id="subjectSelect">
                        <option value="">-- בחר מקצוע --</option>
                    </select>
                </div>
                <div>
                    <span class="selector-label">יח"ל:</span>
                    <select id="unitsSelect">
                        <option value="">-- בחר יח"ל --</option>
                    </select>
                </div>
            </div>

            <p id="summary" class="text-xl font-semibold text-slate-700 mt-4"></p>
            <p class="text-sm text-slate-500 mt-1">הנתונים עובדו על ידי מרכז הידע האזורי __REGION_HEB__</p>
        </header>

        <div id="no-data-msg" class="text-center py-16 text-slate-400 text-lg">
            בחר מקצוע ויחידות לימוד כדי להציג את הנתונים
        </div>

        <main id="dashboard" class="card-container" style="display:none;"></main>

        <section id="charts-section" class="mt-12" style="display:none;">
            <div class="analysis-card mb-8">
                <h2 id="chart1-title" class="text-xl font-bold text-center mb-4"></h2>
                <canvas id="authorityChart"></canvas>
                <p id="chart1-note" class="text-sm text-slate-600 mt-4 text-center"></p>
            </div>
            <div class="analysis-card mb-8">
                <h2 id="chart2-title" class="text-xl font-bold text-center mb-4"></h2>
                <canvas id="yearTrendChart"></canvas>
                <p id="chart2-note" class="text-sm text-slate-600 mt-4 text-center"></p>
            </div>
            <div class="analysis-card mb-8">
                <h2 id="chart3-title" class="text-xl font-bold text-center mb-4"></h2>
                <canvas id="schoolChart"></canvas>
                <p class="text-xs text-slate-500 mt-2 text-center">הנתונים מייצגים ממוצע לשנים שבהן קיים מידע (2021-2024)</p>
            </div>
            <div class="analysis-card mb-8">
                <h2 id="chart5-title" class="text-xl font-bold text-center mb-4"></h2>
                <canvas id="authorityScoreChart"></canvas>
            </div>
            <div class="analysis-card mb-8">
                <h2 id="chart6-title" class="text-xl font-bold text-center mb-4"></h2>
                <canvas id="schoolScoreChart"></canvas>
            </div>
            <div class="analysis-card mb-8">
                <h2 id="chart4-title" class="text-xl font-bold text-center mb-4"></h2>
                <div class="chart-container" style="position: relative; height:300px; width:100%; max-width: 300px; margin: auto;">
                    <canvas id="distributionChart"></canvas>
                </div>
                <div id="school-breakdown" class="mt-6 grid grid-cols-1 md:grid-cols-2 gap-6 text-sm"></div>
            </div>
        </section>

        <footer class="mt-12 text-sm text-slate-600 border-t pt-6">
            <h3 class="font-bold text-base text-slate-800">הערות</h3>
            <ul class="list-disc pr-5 mt-2 space-y-1">
                <li>לא קיים מידע על מוסדות עם פחות מ-11 נבחנים ברמת מוסד, מקצוע ומספר יח"ל, מן הטעם שבעיבוד על אוכלוסייה מצומצמת שינוי קטן של מספר מוחלט עשוי לגרום לשינוי גדול בממוצע, דבר שעשוי לגרום להטעיית הציבור.</li>
                <li>נתוני 2021 נכונים ל-22/04/2023</li>
                <li>נתוני 2022 נכונים ל-26/05/2024</li>
                <li>נתוני 2023 נכונים ל-05/07/2024</li>
                <li>נתוני 2024 נכונים ל-18/07/2025</li>
                <li>עודכן ב-08/02/2026</li>
            </ul>
            <div class="mt-4">
                <h3 class="font-bold text-base text-slate-800">מקורות</h3>
                <ul class="list-disc pr-5 mt-2 space-y-1">
                    <li>התנועה לחופש המידע, משרד החינוך, מינהל טכנולוגיות דיגיטליות ומידע, מרכז מידע בחינות בגרות.</li>
                    <li>דשבורד 'פוטנציאל הון אנושי' של SFI, משרד העבודה ורשות החדשנות.</li>
                </ul>
            </div>
        </footer>
    </div>

<script>
// ============================================================
// EMBEDDED DATA
// ============================================================
const rawData = __RAW_DATA__;
const allHighSchools = __ALL_SCHOOLS__;
const nationalData = __NATIONAL_DATA__;

// ============================================================
// CHART STATE
// ============================================================
let activeCharts = {};

function destroyCharts() {
    Object.values(activeCharts).forEach(c => c.destroy());
    activeCharts = {};
}

// ============================================================
// INITIALIZATION
// ============================================================
function init() {
    Chart.defaults.font.family = "'Assistant', sans-serif";
    Chart.defaults.font.size = 14;

    // Core subject patterns for categorization
    const corePatterns = [
        'אנגלית', 'פיסיקה', 'מתמטיקה', 'ביולוגיה', 'כימיה',
        'ספרות', 'אזרחות', 'תנ"ך', 'תנך',
        'היסטוריה', 'הסטוריה',
        'עברית', 'ערבית',
        'לשון'
    ];

    function isCoreSubject(name) {
        return corePatterns.some(p => name.includes(p));
    }

    const subjects = [...new Set(rawData.map(d => d.subject))].sort((a, b) => a.localeCompare(b, 'he'));
    const coreSubjects = subjects.filter(isCoreSubject);
    const additionalSubjects = subjects.filter(s => !isCoreSubject(s));

    const subjectSelect = document.getElementById('subjectSelect');

    if (coreSubjects.length > 0) {
        const coreGroup = document.createElement('optgroup');
        coreGroup.label = 'מקצועות ליבה';
        coreSubjects.forEach(s => {
            const opt = document.createElement('option');
            opt.value = s;
            opt.textContent = s;
            coreGroup.appendChild(opt);
        });
        subjectSelect.appendChild(coreGroup);
    }

    if (additionalSubjects.length > 0) {
        const additionalGroup = document.createElement('optgroup');
        additionalGroup.label = 'מקצועות נוספים';
        additionalSubjects.forEach(s => {
            const opt = document.createElement('option');
            opt.value = s;
            opt.textContent = s;
            additionalGroup.appendChild(opt);
        });
        subjectSelect.appendChild(additionalGroup);
    }

    subjectSelect.addEventListener('change', onSubjectChange);
    document.getElementById('unitsSelect').addEventListener('change', updateDashboard);

    document.getElementById('no-data-msg').style.display = 'block';
}

function onSubjectChange() {
    const subject = document.getElementById('subjectSelect').value;
    const unitsSelect = document.getElementById('unitsSelect');
    unitsSelect.innerHTML = '<option value="">-- בחר יח"ל --</option>';

    if (!subject) {
        showNoData();
        return;
    }

    const units = [...new Set(rawData.filter(d => d.subject === subject).map(d => d.units))].sort((a, b) => a - b);
    units.forEach(u => {
        const opt = document.createElement('option');
        opt.value = u;
        opt.textContent = u + ' יח"ל';
        unitsSelect.appendChild(opt);
    });

    if (units.length === 1) {
        unitsSelect.value = units[0];
        updateDashboard();
    } else {
        showNoData();
    }
}

function showNoData() {
    document.getElementById('no-data-msg').style.display = 'block';
    document.getElementById('dashboard').style.display = 'none';
    document.getElementById('charts-section').style.display = 'none';
    document.getElementById('summary').textContent = '';
    destroyCharts();
}

// ============================================================
// MAIN UPDATE
// ============================================================
function updateDashboard() {
    const subject = document.getElementById('subjectSelect').value;
    const units = parseInt(document.getElementById('unitsSelect').value);
    if (!subject || isNaN(units)) { showNoData(); return; }

    const filtered = rawData.filter(d => d.subject === subject && d.units === units);
    if (filtered.length === 0) { showNoData(); return; }

    document.getElementById('no-data-msg').style.display = 'none';
    document.getElementById('dashboard').style.display = '';
    document.getElementById('charts-section').style.display = '';

    const aggregated = computeAggregated(filtered);
    const latestYear = Math.max(...filtered.map(d => d.year));
    const latestYearData = filtered.filter(d => d.year === latestYear);
    const totalLatestExaminees = latestYearData.reduce((s, d) => s + d.examinees, 0);

    // National average for this subject+units
    const natEntries = nationalData.filter(d => d.subject === subject && d.units === units);
    const latestNatYear = natEntries.length > 0 ? Math.max(...natEntries.map(d => d.year)) : null;
    const latestNat = natEntries.find(d => d.year === latestNatYear);
    const natAvgScore = latestNat ? latestNat.score : null;
    const natTotalExaminees = latestNat ? latestNat.examinees : null;

    document.getElementById('summary').textContent =
        `בשנת ${latestYear} היו ${totalLatestExaminees} נבחנים ב${subject} ${units} יח"ל באשכול __REGION_HEB__` +
        (natTotalExaminees ? ` (ארצי: ${natTotalExaminees.toLocaleString()})` : '');

    renderCards(aggregated, natAvgScore);
    destroyCharts();
    renderCharts(aggregated, filtered, subject, units, latestYear, totalLatestExaminees, natTotalExaminees);
}

// ============================================================
// DATA PROCESSING
// ============================================================
function computeAggregated(filtered) {
    const schoolMap = {};
    filtered.forEach(d => {
        if (!schoolMap[d.school]) {
            schoolMap[d.school] = {
                school: d.school, authority: d.authority,
                supervision: d.supervision, sector: d.sector,
                years: [], scores: [], examinees: []
            };
        }
        const s = schoolMap[d.school];
        s.years.push(d.year);
        s.scores.push(d.score);
        s.examinees.push(d.examinees);
    });

    return Object.values(schoolMap).map(s => {
        const n = s.years.length;
        const avgScore = n > 0 ? +(s.scores.reduce((a, b) => a + b, 0) / n).toFixed(2) : 0;
        const avgExaminees = n > 0 ? +(s.examinees.reduce((a, b) => a + b, 0) / n).toFixed(2) : 0;
        return {
            school: s.school,
            authority: s.authority,
            supervision: s.supervision,
            sector: s.sector,
            years: s.years.sort(),
            avgScore,
            avgExaminees
        };
    });
}

function calculateGroupStats(data) {
    if (data.length === 0) return { avgGrade: '---', avgExaminees: '---', schools: [] };
    const totalAvgExaminees = data.reduce((sum, item) => sum + item.avgExaminees, 0);
    const weightedSum = data.reduce((sum, item) => sum + (item.avgScore * item.avgExaminees), 0);
    const avgGrade = totalAvgExaminees > 0 ? (weightedSum / totalAvgExaminees).toFixed(2) : '---';
    const avgExamineesPerSchool = (totalAvgExaminees / data.length).toFixed(1);
    const schools = data.map(item => ({ name: item.school, years: `(${item.years.join(', ')})` }));
    return { avgGrade, avgExaminees: avgExamineesPerSchool, schools };
}

// ============================================================
// RENDER SECTOR CARDS
// ============================================================
function renderCards(aggregated, natAvgScore) {
    const druzeData = aggregated.filter(d => d.sector === 'דרוזי');
    const arabData = aggregated.filter(d => d.sector === 'ערבי');
    const jewishStateData = aggregated.filter(d => d.sector === 'יהודי' && d.supervision === 'ממלכתי');
    const jewishReligiousData = aggregated.filter(d => d.sector === 'יהודי' && d.supervision === 'ממלכתי דתי');

    const druzeStats = calculateGroupStats(druzeData);
    const arabStats = calculateGroupStats(arabData);
    const jewishStateStats = calculateGroupStats(jewishStateData);
    const jewishReligiousStats = calculateGroupStats(jewishReligiousData);

    const natLabel = natAvgScore ? `ארצי: ${natAvgScore}` : '';

    function createCard(title, stats, colorClasses, isSubCard) {
        const schoolList = stats.schools.map(s =>
            `<li class="text-sm">${s.name} <span class="text-slate-500 text-xs">${s.years}</span></li>`
        ).join('');
        const p = isSubCard ? 'p-6' : 'p-8';
        const titleSz = isSubCard ? 'text-xl' : 'text-2xl';
        const statSz = isSubCard ? 'text-2xl' : 'text-3xl';

        return `
            <div class="bg-white rounded-xl shadow-md overflow-hidden ${p} flex flex-col border-t-4 ${colorClasses.border}">
                <h2 class="font-bold ${titleSz} ${colorClasses.text} mb-4">${title}</h2>
                <div class="flex-grow space-y-4">
                    <div class="grid grid-cols-2 gap-4 text-center">
                        <div>
                            <p class="text-sm text-slate-500">ציון ממוצע</p>
                            <p class="${statSz} font-bold">${stats.avgGrade}</p>
                            ${natLabel ? `<p class="text-xs text-slate-500 mt-1">${natLabel}</p>` : ''}
                        </div>
                        <div>
                            <p class="text-sm text-slate-500">ממוצע נבחנים לבי"ס (שנתי)</p>
                            <p class="${statSz} font-bold">${stats.avgExaminees}</p>
                        </div>
                    </div>
                    <div class="pt-4">
                        <h3 class="font-semibold text-slate-600 mb-2">בתי ספר בקבוצה:</h3>
                        <ul class="list-disc pr-5 space-y-1 text-slate-700">${schoolList}</ul>
                    </div>
                </div>
            </div>`;
    }

    const dashboard = document.getElementById('dashboard');

    const hasJewishState = jewishStateData.length > 0;
    const hasJewishReligious = jewishReligiousData.length > 0;
    const hasDruze = druzeData.length > 0;
    const hasArab = arabData.length > 0;

    let html = '';

    if (hasJewishState || hasJewishReligious) {
        const jsCard = hasJewishState ? createCard('ממלכתי', jewishStateStats, {border:'border-sky-500', text:'text-sky-600'}, true) : '';
        const jrCard = hasJewishReligious ? createCard('ממלכתי-דתי', jewishReligiousStats, {border:'border-teal-500', text:'text-teal-600'}, true) : '';
        html += `
            <div class="bg-white rounded-xl shadow-md overflow-hidden p-8 flex flex-col border-t-4 border-blue-600 jewish-grid-container">
                <h2 class="font-bold text-2xl text-blue-700 mb-4">המגזר היהודי</h2>
                <div class="jewish-grid flex-grow">${jsCard}${jrCard}</div>
            </div>`;
    }
    if (hasDruze) html += createCard('מגזר דרוזי', druzeStats, {border:'border-violet-500', text:'text-violet-600'}, false);
    if (hasArab) html += createCard('מגזר ערבי', arabStats, {border:'border-amber-500', text:'text-amber-600'}, false);

    if (!html) html = '<p class="text-center text-slate-400 col-span-full py-8">אין נתונים להצגה עבור מקצוע זה</p>';
    dashboard.innerHTML = html;
}

// ============================================================
// RENDER CHARTS
// ============================================================
function renderCharts(aggregated, filtered, subject, units, latestYear, totalLatest, natTotal) {
    // --- Chart 1: Examinees by authority (latest year) ---
    const latestData = filtered.filter(d => d.year === latestYear);
    const authMap = {};
    latestData.forEach(d => { authMap[d.authority] = (authMap[d.authority] || 0) + d.examinees; });
    const authEntries = Object.entries(authMap).sort((a, b) => b[1] - a[1]);

    document.getElementById('chart1-title').textContent = `מספר נבחנים ב${subject} ${units} יח"ל לפי רשות (${latestYear})`;
    document.getElementById('chart1-note').innerHTML =
        `<strong>הערה:</strong> סה"כ ${totalLatest} נבחנים ב${subject} ${units} יח"ל באשכול __REGION_HEB__ בשנת ${latestYear}.` +
        (natTotal ? ` הנתון הארצי: ${natTotal.toLocaleString()}.` : '');

    activeCharts.authority = new Chart(document.getElementById('authorityChart'), {
        type: 'bar',
        data: {
            labels: authEntries.map(e => e[0]),
            datasets: [{
                label: 'מספר נבחנים',
                data: authEntries.map(e => e[1]),
                backgroundColor: 'rgba(22, 163, 74, 0.7)',
                borderColor: 'rgba(22, 163, 74, 1)',
                borderWidth: 1
            }]
        },
        options: {
            responsive: true, indexAxis: 'y',
            scales: {
                x: { beginAtZero: true, ticks: { font: { size: 14 } } },
                y: { ticks: { font: { size: 14 } } }
            },
            plugins: {
                legend: { display: false },
                tooltip: {
                    titleFont: { size: 14 }, bodyFont: { size: 14 },
                    callbacks: { label: ctx => `מספר נבחנים: ${ctx.parsed.x}` }
                }
            }
        }
    });

    // --- Chart 2: Year trend (total examinees per year in cluster) ---
    const yearMap = {};
    const yearScoreMap = {};
    filtered.forEach(d => {
        yearMap[d.year] = (yearMap[d.year] || 0) + d.examinees;
        if (!yearScoreMap[d.year]) yearScoreMap[d.year] = { weightedSum: 0, total: 0 };
        yearScoreMap[d.year].weightedSum += d.score * d.examinees;
        yearScoreMap[d.year].total += d.examinees;
    });
    const years = Object.keys(yearMap).map(Number).sort();

    // National data for note
    const natYearMap = {};
    nationalData.filter(d => d.subject === subject && d.units === units).forEach(d => { natYearMap[d.year] = d.examinees; });

    document.getElementById('chart2-title').textContent = `מגמת מספר נבחנים ב${subject} ${units} יח"ל באשכול לאורך השנים`;
    activeCharts.yearTrend = new Chart(document.getElementById('yearTrendChart'), {
        type: 'bar',
        data: {
            labels: years,
            datasets: [{
                label: 'מספר נבחנים באשכול',
                data: years.map(y => yearMap[y] || 0),
                backgroundColor: 'rgba(75, 192, 192, 0.7)',
                borderColor: 'rgba(75, 192, 192, 1)',
                borderWidth: 1
            }]
        },
        options: {
            responsive: true,
            scales: {
                y: { beginAtZero: true, ticks: { font: { size: 14 } } },
                x: { ticks: { font: { size: 14 } } }
            },
            plugins: {
                legend: { display: false },
                tooltip: {
                    titleFont: { size: 14 }, bodyFont: { size: 14 },
                    callbacks: {
                        afterLabel: function(ctx) {
                            const y = years[ctx.dataIndex];
                            const avg = yearScoreMap[y] ? (yearScoreMap[y].weightedSum / yearScoreMap[y].total).toFixed(2) : '---';
                            let tip = `ציון ממוצע משוקלל: ${avg}`;
                            if (natYearMap[y]) tip += `\nנבחנים ארצי: ${natYearMap[y].toLocaleString()}`;
                            return tip;
                        }
                    }
                }
            }
        }
    });

    // Chart 2 note with national comparison
    const natNotes = years.map(y => natYearMap[y] ? `${y}: ${natYearMap[y].toLocaleString()} (ארצי)` : null).filter(Boolean);
    document.getElementById('chart2-note').innerHTML = natNotes.length > 0
        ? `<strong>השוואה ארצית:</strong> ` + natNotes.join(' | ')
        : '';

    // --- Chart 3: School examinees comparison ---
    const sorted = [...aggregated].sort((a, b) => b.avgExaminees - a.avgExaminees);
    const totalExaminees = aggregated.reduce((s, d) => s + d.avgExaminees, 0);
    const avgExaminees = aggregated.length > 0 ? totalExaminees / aggregated.length : 0;

    document.getElementById('chart3-title').textContent = `השוואת מספר נבחנים ממוצע ב${subject} בין מוסדות הלימוד`;
    activeCharts.school = new Chart(document.getElementById('schoolChart'), {
        type: 'bar',
        data: {
            labels: sorted.map(s => `${s.school} (${s.authority})`),
            datasets: [{
                label: 'מספר נבחנים ממוצע',
                data: sorted.map(s => s.avgExaminees),
                backgroundColor: 'rgba(139, 92, 246, 0.7)',
                borderColor: 'rgba(139, 92, 246, 1)',
                borderWidth: 1
            }]
        },
        options: {
            responsive: true, indexAxis: 'y',
            scales: {
                x: { beginAtZero: true, ticks: { font: { size: 14 } } },
                y: { ticks: { font: { size: 13 } } }
            },
            plugins: {
                legend: { display: false },
                tooltip: {
                    titleFont: { size: 14 }, bodyFont: { size: 14 },
                    callbacks: { label: ctx => `מספר נבחנים ממוצע: ${ctx.parsed.x.toFixed(1)}` }
                },
                annotation: {
                    annotations: {
                        localAvg: {
                            type: 'line',
                            xMin: avgExaminees, xMax: avgExaminees,
                            borderColor: 'rgb(255, 99, 132)', borderWidth: 2, borderDash: [6, 6],
                            label: {
                                content: `ממוצע: ${avgExaminees.toFixed(1)}`,
                                position: 'end',
                                backgroundColor: 'rgba(255, 99, 132, 0.8)',
                                font: { family: 'Assistant' },
                                display: true
                            }
                        }
                    }
                }
            }
        }
    });

    // --- Chart 5: Average score by authority ---
    const authScoreMap = {};
    aggregated.forEach(d => {
        if (!authScoreMap[d.authority]) authScoreMap[d.authority] = { weightedSum: 0, totalExaminees: 0 };
        authScoreMap[d.authority].weightedSum += d.avgScore * d.avgExaminees;
        authScoreMap[d.authority].totalExaminees += d.avgExaminees;
    });
    const authScoreEntries = Object.entries(authScoreMap).map(([auth, v]) => ({
        auth, avgScore: v.totalExaminees > 0 ? +(v.weightedSum / v.totalExaminees).toFixed(2) : 0
    })).sort((a, b) => b.avgScore - a.avgScore);
    const overallAvgScore = (() => {
        const tot = aggregated.reduce((s, d) => s + d.avgExaminees, 0);
        return tot > 0 ? aggregated.reduce((s, d) => s + d.avgScore * d.avgExaminees, 0) / tot : 0;
    })();

    document.getElementById('chart5-title').textContent = `ציון ממוצע ב${subject} ${units} יח"ל לפי רשות מקומית`;
    activeCharts.authorityScore = new Chart(document.getElementById('authorityScoreChart'), {
        type: 'bar',
        data: {
            labels: authScoreEntries.map(e => e.auth),
            datasets: [{
                label: 'ציון ממוצע',
                data: authScoreEntries.map(e => e.avgScore),
                backgroundColor: 'rgba(234, 88, 12, 0.7)',
                borderColor: 'rgba(234, 88, 12, 1)',
                borderWidth: 1
            }]
        },
        options: {
            responsive: true, indexAxis: 'y',
            scales: {
                x: { beginAtZero: false, min: Math.max(0, Math.floor(Math.min(...authScoreEntries.map(e => e.avgScore)) - 5)), ticks: { font: { size: 14 } } },
                y: { ticks: { font: { size: 14 } } }
            },
            plugins: {
                legend: { display: false },
                tooltip: {
                    titleFont: { size: 14 }, bodyFont: { size: 14 },
                    callbacks: { label: ctx => `ציון ממוצע: ${ctx.parsed.x}` }
                },
                annotation: {
                    annotations: {
                        clusterAvg: {
                            type: 'line',
                            xMin: overallAvgScore, xMax: overallAvgScore,
                            borderColor: 'rgb(255, 99, 132)', borderWidth: 2, borderDash: [6, 6],
                            label: {
                                content: `ממוצע אשכול: ${overallAvgScore.toFixed(2)}`,
                                position: 'end',
                                backgroundColor: 'rgba(255, 99, 132, 0.8)',
                                font: { family: 'Assistant', size: 13 },
                                display: true
                            }
                        }
                    }
                }
            }
        }
    });

    // --- Chart 6: Average score by school ---
    const schoolScoreSorted = [...aggregated].sort((a, b) => b.avgScore - a.avgScore);

    document.getElementById('chart6-title').textContent = `ציון ממוצע ב${subject} ${units} יח"ל לפי מוסד לימוד`;
    activeCharts.schoolScore = new Chart(document.getElementById('schoolScoreChart'), {
        type: 'bar',
        data: {
            labels: schoolScoreSorted.map(s => `${s.school} (${s.authority})`),
            datasets: [{
                label: 'ציון ממוצע',
                data: schoolScoreSorted.map(s => s.avgScore),
                backgroundColor: 'rgba(16, 185, 129, 0.7)',
                borderColor: 'rgba(16, 185, 129, 1)',
                borderWidth: 1
            }]
        },
        options: {
            responsive: true, indexAxis: 'y',
            scales: {
                x: { beginAtZero: false, min: Math.max(0, Math.floor(Math.min(...schoolScoreSorted.map(s => s.avgScore)) - 5)), ticks: { font: { size: 14 } } },
                y: { ticks: { font: { size: 13 } } }
            },
            plugins: {
                legend: { display: false },
                tooltip: {
                    titleFont: { size: 14 }, bodyFont: { size: 14 },
                    callbacks: {
                        label: ctx => `ציון ממוצע: ${ctx.parsed.x}`,
                        afterLabel: ctx => {
                            const s = schoolScoreSorted[ctx.dataIndex];
                            return `שנים: ${s.years.join(', ')}\nנבחנים ממוצע: ${s.avgExaminees}`;
                        }
                    }
                },
                annotation: {
                    annotations: {
                        clusterAvg: {
                            type: 'line',
                            xMin: overallAvgScore, xMax: overallAvgScore,
                            borderColor: 'rgb(255, 99, 132)', borderWidth: 2, borderDash: [6, 6],
                            label: {
                                content: `ממוצע אשכול: ${overallAvgScore.toFixed(2)}`,
                                position: 'end',
                                backgroundColor: 'rgba(255, 99, 132, 0.8)',
                                font: { family: 'Assistant', size: 13 },
                                display: true
                            }
                        }
                    }
                }
            }
        }
    });

    // --- Chart 4: Schools with/without this subject ---
    const schoolsWithSubject = [...new Set(aggregated.map(s => s.school))];
    const withCount = schoolsWithSubject.length;
    const withoutCount = allHighSchools.length - withCount;

    document.getElementById('chart4-title').textContent = `התפלגות תיכונים לפי הגשה לבגרות ב${subject}`;
    activeCharts.distribution = new Chart(document.getElementById('distributionChart'), {
        type: 'doughnut',
        data: {
            labels: [`מגישים לבגרות (${withCount})`, `לא מגישים / פחות מ-11 נבחנים (${withoutCount})`],
            datasets: [{
                label: 'מספר תיכונים',
                data: [withCount, withoutCount],
                backgroundColor: ['rgba(59, 130, 246, 0.7)', 'rgba(203, 213, 225, 0.7)'],
                borderColor: ['rgba(59, 130, 246, 1)', 'rgba(203, 213, 225, 1)'],
                borderWidth: 1
            }]
        },
        options: {
            responsive: true, maintainAspectRatio: false,
            plugins: {
                legend: { position: 'top', labels: { font: { size: 14 } } },
                tooltip: {
                    titleFont: { size: 14 }, bodyFont: { size: 14 },
                    callbacks: { label: ctx => `${ctx.label}: ${ctx.raw}` }
                }
            }
        }
    });

    // School breakdown lists
    const breakdown = document.getElementById('school-breakdown');
    const schoolsWithoutSubject = allHighSchools.filter(s => !schoolsWithSubject.includes(s));
    const withList = schoolsWithSubject.sort((a, b) => a.localeCompare(b, 'he')).map(s => `<li>${s}</li>`).join('');
    const withoutList = schoolsWithoutSubject.sort((a, b) => a.localeCompare(b, 'he')).map(s => `<li>${s}</li>`).join('');

    breakdown.innerHTML = `
        <div>
            <h3 class="font-semibold text-slate-700 mb-2 border-b pb-1">מגישים לבגרות ב${subject} (${schoolsWithSubject.length})</h3>
            <ul class="list-disc pr-5 space-y-1 text-slate-600">${withList}</ul>
        </div>
        <div>
            <h3 class="font-semibold text-slate-700 mb-2 border-b pb-1">לא מגישים / פחות מ-11 נבחנים (${schoolsWithoutSubject.length})</h3>
            <ul class="list-disc pr-5 space-y-1 text-slate-600">${withoutList}</ul>
        </div>
    `;
}

// ============================================================
// START
// ============================================================
init();
</script>
</body>
</html>
'''


def main():
    output_dir = os.path.join(BASE_DIR, 'raw_data')
    os.makedirs(output_dir, exist_ok=True)

    print("Reading national data (sheet 1)...")
    national_data = read_national_data()
    print(f"  National data rows: {len(national_data)}")

    for region_name, authorities, all_schools, html_filename, xlsx_filename in [
        ('גליל מערבי', WEST_GALILEE_AUTHORITIES, WEST_ALL_SCHOOLS,
         'west_galilee_dashboard.html', 'west_galilee_raw_data.xlsx'),
        ('גליל מזרחי', EAST_GALILEE_AUTHORITIES, EAST_ALL_SCHOOLS,
         'east_galilee_dashboard.html', 'east_galilee_raw_data.xlsx'),
    ]:
        print(f"\n{'='*60}")
        print(f"Processing: {region_name}")
        print(f"{'='*60}")

        data = read_school_data(authorities)

        # Generate HTML dashboard
        generate_html(
            region_name, data, all_schools, national_data,
            os.path.join(BASE_DIR, html_filename)
        )

        # Generate raw XLSX for validation
        generate_raw_xlsx(data, os.path.join(output_dir, xlsx_filename))

    print(f"\nDone!")


if __name__ == '__main__':
    main()
